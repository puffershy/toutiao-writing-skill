#!/usr/bin/env python3
"""
Generate Toutiao article .docx:
1. Read article_data.json
2. Baidu + Bing dual-engine image search -> download -> WebP to JPEG
3. Gaussian blur on bottom watermark area
4. Embed into .docx -> archive to Desktop/今日头条/{year}/{month}/{date}/{title}.docx
"""

import json
import os
import re
import sys
import time
import io
import shutil
from pathlib import Path

# Fix Windows encoding
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from PIL import Image, ImageFilter
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from icrawler.builtin import BaiduImageCrawler, BingImageCrawler


# ---------- config ----------
DESKTOP = Path.home() / "Desktop"
OUTPUT_BASE = DESKTOP / "今日头条"
SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_DIR = SCRIPT_DIR.parent
JSON_PATH = SKILL_DIR / "output" / "article_data.json"
TEMP_DIR = SKILL_DIR / "output" / "temp_images"


def ensure_dir(path: Path):
    path.mkdir(parents=True, exist_ok=True)


def safe_filename(text: str) -> str:
    return re.sub(r'[\\/:*?"<>|]', '_', text)


def detect_watermark_region(img: Image.Image) -> bool:
    """Check bottom 15% for watermark-like text (high edge density + brightness)."""
    w, h = img.size
    bottom = img.crop((0, int(h * 0.85), w, h))
    gray = bottom.convert("L")
    pixels = list(gray.getdata())

    edge_density = 0
    bw, bh = bottom.size
    for y in range(bh):
        for x in range(1, bw):
            diff = abs(pixels[y * bw + x] - pixels[y * bw + x - 1])
            if diff > 30:
                edge_density += 1
    edge_ratio = edge_density / (bw * bh)
    avg_brightness = sum(pixels) / len(pixels)

    is_watermark = edge_ratio > 0.03 and avg_brightness > 120
    print(f"    watermark: edge={edge_ratio:.4f} bright={avg_brightness:.0f} -> {'BLUR' if is_watermark else 'skip'}")
    return is_watermark


def blur_watermark(img: Image.Image) -> Image.Image:
    """Gaussian blur (radius=8) on bottom 15% if watermark detected."""
    if not detect_watermark_region(img):
        return img

    w, h = img.size
    bottom_h = int(h * 0.15)
    top_part = img.crop((0, 0, w, h - bottom_h))
    bottom_part = img.crop((0, h - bottom_h, w, h))
    blurred = bottom_part.filter(ImageFilter.GaussianBlur(radius=8))
    result = Image.new("RGB", (w, h))
    result.paste(top_part, (0, 0))
    result.paste(blurred, (0, h - bottom_h))
    return result


def _crawl_with(crawler_class, keyword: str, dl_dir: Path, max_num: int = 5) -> bool:
    """Run a crawler, return True if any images downloaded."""
    # Clean first
    for f in dl_dir.iterdir():
        f.unlink()
    try:
        crawler = crawler_class(
            downloader_threads=3,
            storage={"root_dir": str(dl_dir)},
            log_level=50,
        )
        crawler.crawl(keyword=keyword, max_num=max_num, min_size=(200, 200))
    except Exception as e:
        print(f"      {crawler_class.__name__}: {e}")
        return False

    # Check for valid images
    for f in dl_dir.iterdir():
        if f.suffix.lower() in (".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp"):
            if f.stat().st_size > 2048:
                return True
    return False


def _pick_best_image(dl_dir: Path, idx: int, temp_dir: Path) -> Path | None:
    """Pick the largest valid image from download dir, convert to JPEG."""
    best = None
    best_size = 0
    for f in dl_dir.iterdir():
        if f.suffix.lower() in (".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp"):
            sz = f.stat().st_size
            if sz > best_size and sz > 2048:
                best_size = sz
                best = f

    if best is None:
        return None

    try:
        img = Image.open(best).convert("RGB")
        jpg_path = temp_dir / f"img_{idx:02d}.jpg"
        img.save(jpg_path, "JPEG", quality=90)
        print(f"      -> {jpg_path.name} ({jpg_path.stat().st_size} bytes)")
        shutil.rmtree(dl_dir, ignore_errors=True)
        return jpg_path
    except Exception as e:
        print(f"      convert error: {e}")
        return None


def fetch_image_dual(query: str, idx: int, temp_dir: Path) -> Path | None:
    """
    Try Baidu first (better Chinese results), fall back to Bing.
    """
    dl_dir = temp_dir / f"search_{idx:02d}"
    ensure_dir(dl_dir)

    # ---- Try Baidu first ----
    print(f"    [Baidu] {query}")
    if _crawl_with(BaiduImageCrawler, query, dl_dir, max_num=5):
        result = _pick_best_image(dl_dir, idx, temp_dir)
        if result:
            return result

    # ---- Fallback: Bing ----
    print(f"    [Bing]  {query}")
    if _crawl_with(BingImageCrawler, query, dl_dir, max_num=5):
        result = _pick_best_image(dl_dir, idx, temp_dir)
        if result:
            return result

    shutil.rmtree(dl_dir, ignore_errors=True)
    return None


def write_docx(json_data: dict, image_paths: dict, output_path: Path):
    """Write article into .docx with proper formatting."""
    doc = Document()

    # Default font
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # Title
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(json_data["title"])
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.name = "Microsoft YaHei"
    title_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    doc.add_paragraph()

    # Content nodes
    image_seq = 0
    total_images = sum(1 for n in json_data["content"] if n["type"] == "image")
    for node in json_data["content"]:
        t = node["type"]

        if t == "body":
            para = doc.add_paragraph()
            para.paragraph_format.line_spacing = 1.5
            para.paragraph_format.space_after = Pt(6)
            _add_bold_text(para, node["text"])

        elif t == "section_title":
            para = doc.add_paragraph()
            para.paragraph_format.space_before = Pt(12)
            para.paragraph_format.space_after = Pt(6)
            run = para.add_run(node["text"])
            run.bold = True
            run.font.size = Pt(14)
            run.font.name = "Microsoft YaHei"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

        elif t == "separator":
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run("---")
            run.font.color.rgb = RGBColor(180, 180, 180)
            run.font.size = Pt(10)

        elif t == "image":
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(8)
            para.paragraph_format.space_after = Pt(8)

            img_path = image_paths.get(image_seq)
            image_seq += 1

            if img_path and img_path.exists():
                try:
                    img = Image.open(img_path)
                    img = blur_watermark(img)
                    blurred_path = img_path.parent / f"{img_path.stem}_blur{img_path.suffix}"
                    img.save(blurred_path, quality=90)
                    run = para.add_run()
                    run.add_picture(str(blurred_path), width=Inches(4.5))
                except Exception as e:
                    print(f"    insert error: {e}")
                    _add_placeholder(para, node)
            else:
                _add_placeholder(para, node)

    ensure_dir(output_path.parent)
    doc.save(str(output_path))
    print(f"\nArticle saved: {output_path}")
    print(f"  Images: {len(image_paths)}/{total_images} embedded, {total_images - len(image_paths)} placeholders")


def _add_placeholder(para, node):
    """Blue placeholder text for missing images."""
    desc = node.get("description", node.get("query", ""))
    run = para.add_run(f"[图片：{desc}]")
    run.font.color.rgb = RGBColor(0, 100, 200)
    run.font.size = Pt(9)


def _add_bold_text(para, text: str):
    """Split **bold** markers into bold/normal runs."""
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = para.add_run(part[2:-2])
            run.bold = True
        else:
            run = para.add_run(part)
        run.font.size = Pt(11)
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


# ---------- main ----------
def main():
    if not JSON_PATH.exists():
        print(f"Error: {JSON_PATH} not found")
        sys.exit(1)

    with open(JSON_PATH, "r", encoding="utf-8") as f:
        data = json.load(f)

    title = data["title"]
    date_str = data.get("date", time.strftime("%Y-%m-%d"))
    year, month, day = date_str[:4], date_str[5:7], date_str[8:10]
    safe_title = safe_filename(title)
    output_path = OUTPUT_BASE / year / month / day / f"{safe_title}.docx"

    print(f"Title: {title}")
    print(f"Date:  {date_str}")
    print(f"Output: {output_path}")
    print()

    images = [n for n in data["content"] if n["type"] == "image"]
    print(f"Total images: {len(images)}\n")

    # Clean temp dir before each run to avoid stale image mixing
    if TEMP_DIR.exists():
        shutil.rmtree(TEMP_DIR, ignore_errors=True)
    ensure_dir(TEMP_DIR)
    image_paths = {}

    for seq, node in enumerate(images):
        query = node.get("query", "")
        desc = node.get("description", "")
        print(f"--- Image {seq + 1}/{len(images)}: {desc[:60] if desc else query[:60]}")

        img_path = fetch_image_dual(query, seq, TEMP_DIR)
        if img_path:
            image_paths[seq] = img_path
        else:
            print(f"    -> FALLBACK: placeholder text")

    succeeded = len(image_paths)
    print(f"\nResult: {succeeded}/{len(images)} downloaded\n")

    write_docx(data, image_paths, output_path)
    print(f"Temp dir: {TEMP_DIR}")


if __name__ == "__main__":
    main()
